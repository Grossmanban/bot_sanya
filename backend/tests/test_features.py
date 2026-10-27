import asyncio
import io
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app import app_settings, day_actual, export, jobs, lifecycle, reports
from app import services as svc
from app.config import settings
from app.db import connect
from app.notify import Notifier

TODAY = date(2026, 10, 14)  # Wednesday
THU = TODAY + timedelta(days=1)
P = {"start_time": "16:00", "end_time": "19:00", "volume": 25, "location": "Европа"}


class FakeNotifier(Notifier):
    def __init__(self, conn):
        super().__init__(None, conn)
        self.sent: list[tuple[int, str]] = []
        self.docs: list[tuple[int, str]] = []
        self.photos: list[tuple[int, str]] = []
        self.markups: list = []

    async def send(self, chat_id, text, markup=None):
        self.sent.append((chat_id, text))
        return True

    async def send_document(self, chat_id, data, filename, caption=None):
        self.docs.append((chat_id, filename))
        return True

    async def _deliver(self, chat_id, body, media, markup, kind="photo"):
        self.sent.append((chat_id, body))
        self.markups.append((chat_id, markup))
        self.kinds = [*getattr(self, "kinds", []), kind]
        return 1000 + len(self.sent)

    async def send_photo(self, chat_id, photo, caption, markup=None):
        self.sent.append((chat_id, caption))
        self.photos.append((chat_id, photo))
        return True


def at(d: date, hhmm: str) -> datetime:
    h, m = map(int, hhmm.split(":"))
    return datetime(d.year, d.month, d.day, h, m, tzinfo=settings.tz)


def run(coro):
    return asyncio.run(coro)


def shift(conn, user, day=THU, payload=P):
    return svc.create_shift(conn, user, day, dict(payload), svc.new_key(), today=TODAY).shift


def report(conn, user, s, value, now=None):
    return reports.submit(conn, user, s["id"], value, "photo-file-id", "app", svc.new_key(), now=now or at(THU, "19:30"))


def approved(conn, user, s, value, now=None):
    """Sent and accepted by an owner: every report waits for an admin, also when fact ≥ plan."""
    report(conn, user, s, value, now)
    owner = svc.get_user_by_tg(conn, 999) or svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    return reports.review(conn, owner, s["id"], "approve", svc.new_key())


# ---------------------------------------------------------------- points


def test_points_restrict_location_and_capacity(conn, make_user):
    admin = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b, c = make_user(), make_user("Ваня", "Бургер", None), make_user("Вера", "Котова", "vera")
    # no points configured: free text is fine
    assert shift(conn, a, payload={**P, "location": "Где угодно"})["location"] == "Где угодно"
    svc.create_point(conn, admin, "Европа", 2)
    with pytest.raises(svc.ServiceError) as e:
        svc.create_point(conn, admin, "европа", None)
    assert e.value.code == "point_exists"
    with pytest.raises(svc.ServiceError) as e:
        shift(conn, b, payload={**P, "location": "Неизвестная"})
    assert e.value.code == "unknown_point"
    assert shift(conn, b, payload={**P, "location": "европа"})["location"] == "Европа"  # canonical name
    shift(conn, c, day=THU + timedelta(days=1))
    d = make_user("Лена", "Морозова", "lena")
    shift(conn, d)  # second place on Thursday
    e_user = make_user("Костя", "Новиков", "k")
    with pytest.raises(svc.ServiceError) as e:
        shift(conn, e_user)
    assert e.value.code == "point_full"
    week = svc.week_view(conn, b, 0, today=TODAY)
    assert week["days"][3]["points_load"]["Европа"] == 2


def test_point_rename_moves_upcoming_shifts(conn, make_user):
    admin = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    point = svc.create_point(conn, admin, "Европа", None)
    u = make_user()
    sid = shift(conn, u)["id"]
    svc.update_point(conn, admin, point["id"], {"name": "ТЦ Европа"}, today=TODAY)
    assert svc.get_shift(conn, sid)["location"] == "ТЦ Европа"
    svc.update_point(conn, admin, point["id"], {"active": False})
    assert svc.list_points(conn) == []


# ---------------------------------------------------------------- fact


def test_report_window_auto_approve_and_stats(conn, make_user):
    u = make_user()
    s = shift(conn, u)
    assert reports.deadline(s) == at(THU + timedelta(days=1), "00:00")
    late = shift(conn, u, THU + timedelta(days=1), {**P, "start_time": "20:00", "end_time": "23:30"})
    assert reports.deadline(late) == at(THU + timedelta(days=2), "01:30")  # end + 2h is later than midnight
    with pytest.raises(svc.ServiceError) as e:  # the day before the shift
        report(conn, u, s, 20, now=at(THU - timedelta(days=1), "23:59"))
    assert e.value.code == "too_early"
    # from 00:00 of the shift day — any time, also before the shift starts (16:00)
    assert reports.can_submit(s, at(THU, "00:00")) and reports.view(conn, s, at(THU, "09:00"))["started"]
    with pytest.raises(svc.ServiceError) as e:
        reports.submit(conn, u, s["id"], 20, "", "app", svc.new_key(), now=at(THU, "19:30"))
    assert e.value.code == "photo_required"
    sent = report(conn, u, s, 30)  # fact ≥ plan still waits for an admin (the customer's rule)
    assert sent.report["status"] == "pending_review" and sent.shift["report_state"] == "pending_review"
    res = approved(conn, u, s, 30)  # the employee may replace it while it waits; then an admin accepts
    assert res.changed and res.shift["report_state"] == "approved"
    # with a report in, the plan it was checked against stays: no edits or cancelling by the employee
    for action in (lambda: svc.edit_shift(conn, u, s["id"], {**P, "volume": 20}, svc.new_key(), today=THU),
                   lambda: svc.cancel_shift(conn, u, s["id"], svc.new_key(), today=THU)):
        with pytest.raises(svc.ServiceError) as e:
            action()
        assert e.value.code == "has_report"
    view = svc.serialize_shift(res.shift, THU)
    assert view["fact_pay"] == 30 * 120 and view["report_state"] == "approved"
    with pytest.raises(svc.ServiceError) as e:
        report(conn, u, s, 31)
    assert e.value.code == "report_closed"
    st = svc.stats(conn, u["id"], TODAY, THU)
    assert st["totals"]["fact_volume"] == 30 and st["totals"]["fact_pay"] == 3600


def test_report_review_return_and_fine(conn, make_user):
    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    u = make_user()
    s = shift(conn, u, payload={**P, "volume": 20})
    n = FakeNotifier(conn)
    res = report(conn, u, s, 10)
    assert res.report["status"] == "pending_review" and res.shift["report_state"] == "pending_review"
    # while waiting the employee may replace the report (same attempt)
    res = report(conn, u, s, 12)
    assert res.report["attempt"] == 1 and res.shift["fact_volume"] == 12
    run(n.report_submitted(res.shift, res.report))
    assert any("Отчёт отправлен на проверку" in t for chat, t in n.sent if chat == u["telegram_id"])
    assert (999, "photo-file-id") in n.photos  # reviewers get the screenshot with buttons
    assert reports.queue(conn, "review")[0]["shift_id"] == s["id"]
    assert reports.counts(conn) == {"review": 1, "no_show": 0}

    with pytest.raises(svc.ServiceError) as e:
        reports.review(conn, owner, s["id"], "return", svc.new_key())
    assert e.value.code == "comment_required"
    back = reports.review(conn, owner, s["id"], "return", svc.new_key(), comment="Скрин обрезан")
    assert back.shift["report_state"] == "needs_fix" and back.report["status"] == "rejected_retry"
    # a returned report has no deadline: next day is fine, attempt 2
    again = report(conn, u, s, 14, now=at(THU + timedelta(days=2), "12:00"))
    assert again.report["attempt"] == 2 and again.shift["report_state"] == "pending_review"

    with pytest.raises(svc.ServiceError) as e:  # no fine for a shortfall before the day's global actual
        reports.review(conn, owner, s["id"], "fine", svc.new_key(), units=5, comment="Не выполнена норма")
    assert e.value.code == "shortfall_fine_blocked"
    day_actual.set_actual(conn, owner, THU, 10, svc.new_key(), today=THU + timedelta(days=2))  # team missed the plan
    fined = reports.review(conn, owner, s["id"], "fine", svc.new_key(), units=5, comment="Не выполнена норма")
    assert fined.shift["report_state"] == "rejected" and fined.shift["shortfall_penalty"] == 600
    # the employee's card: units, no rubles; the staff's day view keeps both (ТЗ 27.10 §19)
    card = svc.serialize_shift(fined.shift, THU, conn)
    assert (card["shortfall_units"], card["shortfall_penalty"], card["report"]["fine_units"], card["report"]["fine"]) == (5, None, 5, None)
    assert svc.serialize_shift(fined.shift, THU, conn, staff=True)["report"]["fine"] == 600
    run(n.report_reviewed(fined.shift, fined.report, owner))
    assert any("5 штрафных единиц" in t for chat, t in n.sent if chat == u["telegram_id"])
    assert [r["status"] for r in reports.attempts(conn, s["id"])] == ["rejected_retry", "rejected_penalty"]

    penalties = svc.user_penalties(conn, u["id"])
    assert penalties[0]["kind"] == "report" and penalties[0]["penalty_amount"] == 600
    assert "Не выполнена норма" in penalties[0]["reason"]
    st = svc.stats(conn, u["id"], TODAY, THU)
    assert st["totals"]["penalties"] == 600 and st["totals"]["shortfall_volume"] == 6
    wb = load_workbook(io.BytesIO(export.build_xlsx(conn, TODAY, THU)))
    header = list(wb["Итоги"].iter_rows(values_only=True))[0]
    row = list(wb["Итоги"].iter_rows(values_only=True))[1]
    assert row[header.index("Недобор, шт.")] == 6 and row[header.index("Штрафы, ₽")] == 600


def test_report_permissions(conn, make_user):
    svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    u = make_user()
    trainer = make_user(first="Вера", last="Тренер")
    conn.execute("UPDATE users SET network_id = 1 WHERE id IN (?, ?)", (u["id"], trainer["id"]))
    conn.execute("UPDATE users SET role = 'trainer' WHERE id = ?", (trainer["id"],))
    trainer, u = svc.get_user(conn, trainer["id"]), svc.get_user(conn, u["id"])
    s = shift(conn, u, payload={**P, "volume": 20})
    report(conn, u, s, 10)
    assert trainer["telegram_id"] in reports.reviewers(conn, u)
    ok = reports.review(conn, trainer, s["id"], "approve", svc.new_key())
    assert ok.shift["report_state"] == "approved" and ok.shift["shortfall_status"] == "waived"

    admin = make_user(first="Адм", last="Ин")
    conn.execute("UPDATE users SET role = 'admin' WHERE id = ?", (admin["id"],))
    own = shift(conn, svc.get_user(conn, admin["id"]), payload={**P, "volume": 20})
    report(conn, admin, own, 5)
    with pytest.raises(svc.ServiceError) as e:  # nobody but an owner reviews their own report
        reports.review(conn, svc.get_user(conn, admin["id"]), own["id"], "approve", svc.new_key())
    assert e.value.code == "own_report"


def test_no_show_reminder_and_resolution(conn, make_user):
    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    app_settings.set_raw(conn, "reports_since", TODAY.isoformat())
    u = make_user()
    s = shift(conn, u)
    other = shift(conn, make_user(first="Пётр", last="Сдал"))
    n = FakeNotifier(conn)
    run(jobs.tick(conn, n, at(THU, "21:30")))
    assert not [t for _, t in n.sent if "Вы не сдали отчёт" in t]
    run(jobs.tick(conn, n, at(THU, "22:00")))  # 2 hours before the 00:00 deadline
    reminders = [t for chat, t in n.sent if "Вы не сдали отчёт" in t]
    assert len(reminders) == 2 and "до 00:00" in reminders[0]
    approved(conn, svc.get_user(conn, other["user_id"]), other, 30, now=at(THU, "23:00"))

    digest = [t for chat, t in n.sent if chat == 999 and "Отчёты за" in t]
    assert digest and "Иван Петров" in digest[0]  # 22:00 digest: who hasn't reported

    run(jobs.tick(conn, n, at(THU + timedelta(days=1), "00:01")))
    assert svc.get_shift(conn, s["id"])["report_state"] == "no_show"
    assert svc.get_shift(conn, other["id"])["report_state"] == "approved"
    assert any("неявка" in t for chat, t in n.sent if chat == u["telegram_id"])
    assert any("Неявка" in t for chat, t in n.sent if chat == 999)
    run(jobs.tick(conn, n, at(THU + timedelta(days=1), "00:30")))
    assert len([t for chat, t in n.sent if chat == u["telegram_id"] and "неявка" in t]) == 1
    with pytest.raises(svc.ServiceError) as e:
        report(conn, u, s, 25, now=at(THU + timedelta(days=1), "08:00"))
    assert e.value.code == "deadline_passed"

    assert reports.queue(conn, "no_show")[0]["shift_id"] == s["id"]
    fined = reports.resolve_no_show(conn, owner, s["id"], "fine", svc.new_key(), units=8, comment="Не вышел")
    assert fined.shift["shortfall_penalty"] == 960 and reports.queue(conn, "no_show") == []
    assert svc.user_penalties(conn, u["id"])[0]["reason"].startswith("Неявка")
    # path B: staff enter the result themselves — the fine is cleared, the shift approved
    fixed = reports.manual_result(conn, owner, s["id"], 25, svc.new_key(), comment="Был на точке, скрин в чате")
    assert fixed.shift["report_state"] == "approved" and fixed.shift["shortfall_status"] is None
    assert fixed.report["source"] == "admin"
    assert [p["status"] for p in svc.user_penalties(conn, u["id"])] == ["cancelled"]  # annulled, kept as history


def test_old_shifts_never_become_no_shows(conn, make_user):
    u = make_user()
    s = shift(conn, u)
    n = FakeNotifier(conn)
    run(jobs.tick(conn, n, at(THU + timedelta(days=3), "12:00")))  # first run after the update
    assert svc.get_shift(conn, s["id"])["report_state"] is None and n.sent == []


# ---------------------------------------------------------------- scheduled jobs


def test_eve_reminder_auto_raspred_once(conn, make_user):
    svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b = make_user(), make_user("Ваня", "Бургер", None)
    shift(conn, a)
    shift(conn, b, payload={**P, "start_time": "14:00"})
    b = svc.set_work_status(conn, b, "not_working", svc.new_key())  # «Не работаю» mutes the reminders
    n = FakeNotifier(conn)
    ran = run(jobs.tick(conn, n, at(TODAY, "21:05")))
    assert "eve_reminder" in ran and "auto_raspred" in ran
    reminders = [t for chat, t in n.sent if "Завтра" in t]
    assert len(reminders) == 1 and "На точке также: Ваня 14:00–19:00" in reminders[0]
    assert any(chat == 999 and "Распред на" in t for chat, t in n.sent)
    # second tick the same evening sends nothing again
    n.sent.clear()
    assert run(jobs.tick(conn, n, at(TODAY, "21:20"))) == []
    assert n.sent == []


def test_jobs_respect_window_and_settings(conn, make_user):
    u = make_user()
    shift(conn, u)
    n = FakeNotifier(conn)
    # too early and too late (window is 3h)
    assert "eve_reminder" not in run(jobs.tick(conn, n, at(TODAY, "19:59")))
    app_settings.update(conn, {"eve_reminder_enabled": False})
    assert "eve_reminder" not in run(jobs.tick(conn, n, at(TODAY, "20:10")))
    with pytest.raises(svc.ServiceError):
        app_settings.update(conn, {"eve_reminder_time": "25:00"})
    with pytest.raises(svc.ServiceError):
        app_settings.update(conn, {"nope": True})


def test_week_reminder_only_for_empty_next_week(conn, make_user):
    a, b = make_user(), make_user("Ваня", "Бургер", None)
    next_mon = date(2026, 10, 19)
    shift(conn, b, day=next_mon)
    n = FakeNotifier(conn)
    ran = run(jobs.tick(conn, n, at(date(2026, 10, 15), "12:01")))  # Thursday
    assert "week_reminder" in ran
    targets = {chat for chat, t in n.sent if "следующую неделю" in t}
    assert targets == {a["telegram_id"]}


def test_fact_request_after_shift_end(conn, make_user):
    app_settings.set_raw(conn, "reports_since", TODAY.isoformat())
    u = make_user()
    s = shift(conn, u)
    n = FakeNotifier(conn)
    run(jobs.tick(conn, n, at(THU, "19:10")))
    assert not [t for _, t in n.sent if "скриншот" in t]
    run(jobs.tick(conn, n, at(THU, "19:16")))
    asks = [t for _, t in n.sent if "скриншот" in t]
    assert len(asks) == 1 and "План был 25" in asks[0] and "Сдать до 00:00" in asks[0]
    run(jobs.tick(conn, n, at(THU, "19:30")))
    assert len([t for _, t in n.sent if "скриншот" in t]) == 1
    assert svc.get_shift(conn, s["id"])["fact_asked_at"]


def test_weekly_report_with_excel(conn, make_user):
    svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    u = make_user()
    s = shift(conn, u)
    report(conn, u, s, 22)
    n = FakeNotifier(conn)
    monday = date(2026, 10, 19)
    ran = run(jobs.tick(conn, n, at(monday, "10:00")))
    assert "weekly_report" in ran
    summary = next(t for _, t in n.sent if "Итоги недели" in t)
    assert "план 25, факт 22 шт." in summary
    assert n.docs and n.docs[0][1].endswith(".xlsx")


# ---------------------------------------------------------------- export, broadcasts, migration


def test_excel_export_contents(conn, make_user):
    u = make_user()
    s = shift(conn, u)
    report(conn, u, s, 20)
    wb = load_workbook(io.BytesIO(export.build_xlsx(conn, TODAY, THU)))
    assert wb.sheetnames == ["Итоги", "Смены", "Штрафы и бонусы"]
    totals = list(wb["Итоги"].iter_rows(values_only=True))
    assert totals[1][0] == "Иван Петров" and totals[1][3] == 25 and totals[1][4] == 20
    assert totals[-1][0] == "Итого"
    assert list(wb["Смены"].iter_rows(values_only=True))[1][5] == "Европа"


def test_broadcast_audiences(conn, make_user):
    admin = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    make_user()
    make_user("Ваня", "Бургер", None, working=False)
    assert len(svc.broadcast_recipients(conn, "working")) == 1
    assert len(svc.broadcast_recipients(conn, "all")) == 3  # owner included
    row, recipients = svc.create_broadcast(conn, admin, "all", "  Привет всем  ")
    assert row["text"] == "Привет всем" and row["recipients"] == 3
    n = FakeNotifier(conn)
    assert run(n.broadcast(row["id"], recipients, row["html"])) == 3
    assert svc.list_broadcasts(conn)[0]["delivered"] == 3
    with pytest.raises(svc.ServiceError):
        svc.create_broadcast(conn, admin, "all", "   ")


def test_migration_from_first_schema(tmp_path):
    import sqlite3

    from app.db import SCHEMA

    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.executescript(SCHEMA)
    old.execute("INSERT INTO users (telegram_id, first_name, last_name, created_at) VALUES (1, 'А', 'Б', 'x')")
    old.commit()
    old.close()
    conn = connect(str(path))
    from app.db import MIGRATIONS

    assert conn.execute("PRAGMA user_version").fetchone()[0] == len(MIGRATIONS)
    assert conn.execute("SELECT notify_reminders FROM users").fetchone()[0] == 0  # «Не работаю» → no reminders
    connect(str(path))  # idempotent


def test_broadcast_audience_toggles_union(conn, make_user):
    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")  # admin, not working
    worker = make_user()
    idle = make_user("Ваня", "Бургер", None, working=False)
    helper = make_user("Мария", "Админова", "m")  # working
    svc.set_role(conn, owner, helper["id"], "admin", svc.new_key())
    ids = lambda groups: {u["id"] for u in svc.broadcast_recipients(conn, groups)}
    assert ids(["working"]) == {worker["id"], helper["id"]}
    assert ids(["not_working"]) == {owner["id"], idle["id"]}
    assert ids(["admins"]) == {owner["id"], helper["id"]}
    assert ids(["working", "admins"]) == {worker["id"], helper["id"], owner["id"]}  # helper counted once
    with pytest.raises(svc.ServiceError):
        svc.broadcast_recipients(conn, [])
    row, rec = svc.create_broadcast(conn, owner, ["working", "admins"], "<b>Важно</b>: завтра <i>сбор</i> в 9:00", signed=False)
    assert row["html"] == "<b>Важно</b>: завтра <i>сбор</i> в 9:00" and row["text"].startswith("Важно")
    assert row["signed"] == 0 and len(rec) == 3


def test_sanitizer_keeps_telegram_tags_and_drops_the_rest():
    from app.richtext import sanitize

    html, text = sanitize(
        '<div><strong>Жирный</strong> <em>курсив</em> <u>под</u> <strike>зач</strike></div>'
        '<div><span class="tg-spoiler">секрет</span> <code>код <b>без</b> формата</code></div>'
        '<blockquote>цитата <blockquote>вложенная</blockquote></blockquote>'
        '<a href="https://t.me/x">ссылка</a> <a href="javascript:alert(1)">плохая</a>'
        '<script>alert(1)</script><img src=x onerror=alert(1)><span style="color:red">цвет</span> 5 < 6 & 7'
    )
    assert "<b>Жирный</b>" in html and "<i>курсив</i>" in html and "<u>под</u>" in html and "<s>зач</s>" in html
    assert "<tg-spoiler>секрет</tg-spoiler>" in html
    assert "<code>код без формата</code>" in html
    assert html.count("<blockquote>") == 1
    assert '<a href="https://t.me/x">ссылка</a>' in html and "javascript" not in html and "плохая" in html
    assert "<script" not in html and "<img" not in html and "style" not in html
    assert "5 &lt; 6 &amp; 7" in html
    assert "Жирный" in text and "<" not in text.replace("5 < 6", "")


def test_sanitizer_closes_open_tags_and_rejects_empty():
    from app.richtext import sanitize

    html, _ = sanitize("<b>начало <i>середина")
    assert html == "<b>начало <i>середина</i></b>"
    html, _ = sanitize("<b>a<i>b</b>c</i>")
    assert html == "<b>a<i>b</i></b>c"
    with pytest.raises(svc.ServiceError):
        svc.prepare_message("<b> </b><br>")
    with pytest.raises(svc.ServiceError):
        svc.prepare_message("x" * 4001)


# ---------------------------------------------------------------- plan vs fact shortfall


def test_archive_posts_with_hashtags_and_retries(conn, make_user):
    from aiogram.exceptions import TelegramRetryAfter

    notifier = FakeNotifier(conn)
    user = make_user(first="Олег", last="Сидоров-Петров")
    assert not run(notifier.archive(["штраф"], "x", user))  # no group yet
    app_settings.set_raw(conn, "archive_chat_id", "-1001")
    assert run(notifier.archive(["штраф", "недобор"], "💸 500 ₽", user))
    chat, text = notifier.sent[-1]
    assert chat == -1001 and text.startswith("#штраф #недобор #Олег_Сидоров_Петров")
    assert conn.execute("SELECT sent FROM archive_log ORDER BY id DESC").fetchone()[0] == 1

    calls = []

    async def flaky():
        calls.append(1)
        if len(calls) < 2:
            raise TelegramRetryAfter(method=None, message="flood", retry_after=0)
        return "ok"

    assert run(Notifier._retry(flaky)) == "ok" and len(calls) == 2


# ---------------------------------------------------------------- the day's global actual (ТЗ 3.5)


def test_global_actual_recalculation(conn, make_user):
    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b, c = make_user(first="Аня"), make_user(first="Боря"), make_user(first="Вика")
    sa = shift(conn, a, payload={**P, "volume": 20})
    sb = shift(conn, b, payload={**P, "volume": 20})
    sc = shift(conn, c, payload={**P, "volume": 20})
    assert day_actual.day_plan(conn, THU) == 60
    report(conn, a, sa, 10)  # below plan → review
    report(conn, b, sb, 15)  # below plan → review
    later = THU + timedelta(days=1)

    # the team missed the plan: below-plan reports stay for manual review, fines allowed
    miss = day_actual.set_actual(conn, owner, THU, 50, svc.new_key(), today=later)
    assert miss.changed and not miss.row["met"] and miss.approved == []
    reports.review(conn, owner, sb["id"], "fine", svc.new_key(), units=4, comment="Не выполнена норма")
    assert svc.get_shift(conn, sa["id"])["report_state"] == "pending_review"

    # corrected upwards: the team met it → pending approved, the fine offered for annulment
    met = day_actual.set_actual(conn, owner, THU, 64, svc.new_key(), today=later)
    assert met.row["met"] and [s["id"] for s in met.approved] == [sa["id"]]
    assert reports.latest(conn, sa["id"])["status"] == "approved_team"
    assert svc.get_shift(conn, sa["id"])["shortfall_status"] == "waived"
    assert [s["id"] for s in met.fined] == [sb["id"]]
    assert [h["new_value"] for h in day_actual.history(conn, THU)] == [64, 50]  # every change is logged

    n = FakeNotifier(conn)
    run(n.global_actual_set(THU, met, owner))
    assert any("принят автоматически" in t for chat, t in n.sent if chat == a["telegram_id"])

    # a report sent after the actual on a day the team met the plan is approved at once; fines are blocked
    late = report(conn, c, sc, 5, now=at(THU, "21:00"))
    assert late.report["status"] == "approved_team" and late.shift["report_state"] == "approved"
    with pytest.raises(svc.ServiceError) as e:
        reports.review(conn, owner, sc["id"], "fine", svc.new_key(), units=4, comment="x")
    assert "выполнила" in e.value.message

    done = day_actual.annul_fines(conn, owner, THU, None, svc.new_key())
    assert [s["id"] for s in done] == [sb["id"]] and done[0]["annulled_amount"] == 480
    assert svc.get_shift(conn, sb["id"])["shortfall_penalty"] is None
    assert [p["status"] for p in svc.user_penalties(conn, b["id"])] == ["cancelled"]
    run(n.fines_annulled(done, owner))
    assert any("аннулирован" in t for chat, t in n.sent if chat == b["telegram_id"])


def test_global_actual_is_owner_only_and_not_for_the_future(conn, make_user):
    svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    admin = make_user(first="Адм")
    conn.execute("UPDATE users SET role = 'admin' WHERE id = ?", (admin["id"],))
    shift(conn, make_user())
    with pytest.raises(svc.ServiceError) as e:
        day_actual.set_actual(conn, svc.get_user(conn, admin["id"]), THU, 30, svc.new_key(), today=THU)
    assert e.value.status == 403
    owner = svc.get_user_by_tg(conn, 999)
    with pytest.raises(svc.ServiceError) as e:
        day_actual.set_actual(conn, owner, THU, 30, svc.new_key(), today=TODAY)
    assert e.value.code == "future_day"


def test_global_actual_reminders_and_missing_days(conn, make_user):
    svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    shift(conn, make_user())
    n = FakeNotifier(conn)
    run(jobs.tick(conn, n, at(THU, "21:30")))
    evening = [t for chat, t in n.sent if chat == 999 and "Не введён общий факт за сегодня" in t]
    assert len(evening) == 1 and "План команды 25" in evening[0]
    assert day_actual.missing_days(conn, THU + timedelta(days=1)) == [THU]
    run(jobs.tick(conn, n, at(THU + timedelta(days=1), "09:30")))
    assert any("за вчера" in t for chat, t in n.sent if chat == 999)
    day_actual.set_actual(conn, svc.get_user_by_tg(conn, 999), THU, 25, svc.new_key(), today=THU + timedelta(days=1))
    assert day_actual.missing_days(conn, THU + timedelta(days=1)) == []
    view = day_actual.overview(conn, THU + timedelta(days=1))
    assert view[0]["date"] == THU.isoformat() and view[0]["met"] is True and view[0]["plan"] == 25


def test_migration_keeps_reports_when_adding_team_status(tmp_path):
    import sqlite3

    from app import db

    path = tmp_path / "old.db"
    old = sqlite3.connect(path, isolation_level=None)
    old.row_factory = sqlite3.Row
    old.executescript(db.SCHEMA)
    saved = db.MIGRATIONS[:]
    try:
        db.MIGRATIONS[:] = saved[:6]
        db.migrate(old)
        old.execute("INSERT INTO users (telegram_id, first_name, last_name, status, created_at) VALUES (1,'a','b','approved','x')")
        old.execute("INSERT INTO shifts (user_id, shift_date, start_time, end_time, volume, location, status, created_at, "
                    "updated_at, fact_volume, report_state) VALUES (1,'2026-10-01','10:00','12:00',20,'x','active','t','t',5,"
                    "'pending_review')")
        old.execute("INSERT INTO reports (shift_id, user_id, attempt, fact_volume, photo_file_id, source, status, "
                    "submitted_by, submitted_at) VALUES (1, 1, 1, 5, 'f', 'app', 'pending_review', 1, 't')")
        db.MIGRATIONS[:] = saved
        db.migrate(old)
    finally:
        db.MIGRATIONS[:] = saved
    assert old.execute("PRAGMA user_version").fetchone()[0] == len(saved)
    assert old.execute("SELECT status, photo_file_id FROM reports").fetchone()[:] == ("pending_review", "f")
    old.execute("UPDATE reports SET status = 'approved_team'")  # the new status is accepted


# ---------------------------------------------------------------- money (ТЗ модуль 4)


def _staff(conn, make_user, role, first, network=None):
    u = make_user(first=first)
    conn.execute("UPDATE users SET role = ?, network_id = ? WHERE id = ?", (role, network, u["id"]))
    return svc.get_user(conn, u["id"])


def test_final_cancel_creates_a_draft_fine(conn, make_user):
    from app import finance

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    u = make_user()
    s = shift(conn, u)
    # shifts finally cancelled before 27.10 (no such action any more) keep their draft fine
    conn.execute("UPDATE shifts SET status = 'finally_cancelled' WHERE id = ?", (s["id"],))
    finance.final_cancel_draft(conn, svc.get_shift(conn, s["id"]), u["telegram_id"])
    [draft] = finance.linked(conn, s["id"], ("final_cancel",))
    assert draft["status"] == "draft" and draft["amount"] is None
    assert finance.balance(conn, u["id"])["drafts"] == 1 and finance.balance(conn, u["id"])["balance"] == 0
    svc.set_penalty(conn, s["id"], 1000, 999)  # the bot's preset button
    [fine] = finance.linked(conn, s["id"], ("final_cancel",))
    assert fine["status"] == "active" and fine["amount"] == 1000
    assert finance.balance(conn, u["id"])["balance"] == -1000
    # restoring the shift annuls the fine
    svc.admin_edit_shift(conn, owner, s["id"], {"action": "force_restore", "reason": "ошибка"}, svc.new_key())
    assert finance.linked(conn, s["id"], ("final_cancel",)) == [] and finance.balance(conn, u["id"])["balance"] == 0


def test_trainer_fine_needs_approval_and_balance(conn, make_user):
    from app import finance

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer = _staff(conn, make_user, "trainer", "Тренер", network=1)
    admin = _staff(conn, make_user, "admin", "Админ")
    u = make_user(first="Иван")
    conn.execute("UPDATE users SET network_id = 1 WHERE id = ?", (u["id"],))
    stranger = make_user(first="Чужой")

    req = finance.create(conn, trainer, u["id"], "fine", None, "Опоздал", svc.new_key(), units=4)
    assert req.tx["status"] == "pending_admin_approval" and finance.balance(conn, u["id"])["balance"] == 0
    with pytest.raises(svc.ServiceError):  # not their network
        finance.create(conn, trainer, stranger["id"], "fine", None, "x", svc.new_key(), units=4)
    with pytest.raises(svc.ServiceError) as e:  # owners and admins decide, not trainers
        finance.decide(conn, trainer, req.tx["id"], True, None, svc.new_key())
    assert e.value.status == 403
    ok = finance.decide(conn, admin, req.tx["id"], True, None, svc.new_key())
    assert ok.tx["status"] == "active" and finance.balance(conn, u["id"])["balance"] == -480

    bonus = finance.create(conn, trainer, u["id"], "bonus", 300, "Лучший день", svc.new_key())
    assert bonus.tx["status"] == "active" and finance.balance(conn, u["id"])["balance"] == -180
    with pytest.raises(svc.ServiceError) as e:
        finance.create(conn, admin, u["id"], "deduction", 100, "Форма", svc.new_key())
    assert e.value.code == "owner_only"
    finance.create(conn, owner, u["id"], "deduction", 100, "Форма", svc.new_key())
    assert finance.balance(conn, u["id"]) | {} == {"fines": 480, "bonuses": 300, "deductions": 100, "balance": -280,
                                                     "pending": 0, "drafts": 0}

    declined = finance.create(conn, trainer, u["id"], "fine", None, "Спорно", svc.new_key(), units=5)
    finance.decide(conn, owner, declined.tx["id"], False, "Нет доказательств", svc.new_key())
    assert finance.get(conn, declined.tx["id"])["status"] == "rejected"

    with pytest.raises(svc.ServiceError) as e:  # a reason is required to annul
        finance.cancel(conn, admin, bonus.tx["id"], " ", svc.new_key())
    assert e.value.code == "reason_required"
    finance.cancel(conn, admin, bonus.tx["id"], "Ошибка", svc.new_key())
    with pytest.raises(svc.ServiceError) as e:
        finance.pay_all(conn, admin, u["id"], svc.new_key())
    assert e.value.status == 403
    paid = finance.pay_all(conn, owner, u["id"], svc.new_key())
    assert [p["amount"] for p in paid] == [480]
    assert finance.balance(conn, u["id"])["fines"] == 0 and finance.balance(conn, u["id"])["deductions"] == 100

    people = finance.summary(conn, only={u["id"]})
    assert people[0]["balance"] == -100
    n = FakeNotifier(conn)
    run(n.tx_created(req.tx, trainer))
    assert any("Заявка на штраф" in t for chat, t in n.sent if chat == 999)


def test_trainer_report_fine_waits_and_reject_clears_it(conn, make_user):
    from app import finance

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer = _staff(conn, make_user, "trainer", "Тренер", network=1)
    u = make_user(first="Иван")
    conn.execute("UPDATE users SET network_id = 1 WHERE id = ?", (u["id"],))
    u = svc.get_user(conn, u["id"])
    s = shift(conn, u, payload={**P, "volume": 20})
    report(conn, u, s, 10)
    day_actual.set_actual(conn, owner, THU, 5, svc.new_key(), today=THU)
    res = reports.review(conn, trainer, s["id"], "fine", svc.new_key(), units=4, comment="Не выполнена норма")
    [tx] = finance.linked(conn, s["id"], ("report",))
    assert tx["status"] == "pending_admin_approval" and finance.balance(conn, u["id"])["balance"] == 0
    view = reports.view(conn, res.shift)
    assert view["fine"] is None and view["fine_pending"] == 480 and view["fine_pending_units"] == 4
    finance.decide(conn, owner, tx["id"], False, None, svc.new_key())
    shift_now = svc.get_shift(conn, s["id"])
    assert shift_now["shortfall_status"] == "waived" and shift_now["shortfall_penalty"] is None


def test_migration_moves_fines_into_the_ledger(tmp_path):
    import sqlite3

    from app import db

    old = sqlite3.connect(tmp_path / "old.db", isolation_level=None)
    old.row_factory = sqlite3.Row
    old.executescript(db.SCHEMA)
    saved = db.MIGRATIONS[:]
    try:
        db.MIGRATIONS[:] = saved[:7]
        db.migrate(old)
        old.execute("INSERT INTO users (telegram_id, first_name, last_name, status, created_at) VALUES (1,'a','b','approved','x')")
        ins = ("INSERT INTO shifts (user_id, shift_date, start_time, end_time, volume, location, status, created_at, updated_at,"
               " penalty_amount, fact_volume, shortfall_status, shortfall_penalty, report_state) VALUES (1,?,'10:00','12:00',20,"
               "'x',?,'t','t',?,?,?,?,?)")
        old.execute(ins, ("2026-10-01", "finally_cancelled", 1000, None, None, None, None))
        old.execute(ins, ("2026-10-02", "finally_cancelled", None, None, None, None, None))
        old.execute(ins, ("2026-10-03", "active", None, 10, "fined", 700, "rejected"))
        old.execute(ins, ("2026-10-04", "active", None, None, "fined", 300, "no_show"))
        db.MIGRATIONS[:] = saved
        db.migrate(old)
    finally:
        db.MIGRATIONS[:] = saved
    rows = [tuple(r) for r in old.execute("SELECT source, status, amount FROM financial_transactions ORDER BY shift_id")]
    assert rows == [("final_cancel", "active", 1000), ("final_cancel", "draft", None), ("report", "active", 700),
                    ("no_show", "active", 300)]


# ---------------------------------------------------------------- networks (ТЗ модуль 8)


def test_networks_trainers_and_scope(conn, make_user):
    from app import networks
    from app import permissions as perms

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    admin = _staff(conn, make_user, "admin", "Админ")
    a, b, c = make_user(first="Аня"), make_user(first="Боря"), make_user(first="Вика")
    with pytest.raises(svc.ServiceError) as e:
        networks.create(conn, admin, "Север", svc.new_key())
    assert e.value.status == 403
    north = networks.create(conn, owner, "Север", svc.new_key())
    south = networks.create(conn, owner, "Юг", svc.new_key())
    with pytest.raises(svc.ServiceError) as e:
        networks.create(conn, owner, "  север ", svc.new_key())
    assert e.value.code == "duplicate"

    networks.set_trainer(conn, owner, north["id"], a["id"], svc.new_key())
    a = svc.get_user(conn, a["id"])
    assert a["role"] == "trainer" and a["network_id"] == north["id"]
    networks.set_member(conn, owner, b["id"], north["id"], svc.new_key())
    networks.set_member(conn, owner, c["id"], south["id"], svc.new_key())
    b, c = svc.get_user(conn, b["id"]), svc.get_user(conn, c["id"])
    assert perms.manages(a, b) and not perms.manages(a, c)
    assert networks.trainer_of(conn, b)["id"] == a["id"] and networks.trainer_of(conn, a) is None

    tree = networks.listing(conn, owner)
    assert [n["name"] for n in tree] == ["Север", "Юг"] and tree[0]["trainer"]["id"] == a["id"]
    assert [m["id"] for m in tree[0]["members"]] == [b["id"]]
    assert [n["name"] for n in networks.listing(conn, a)] == ["Север"]  # a trainer sees only their network

    # a new trainer: the previous one stays in the network as an employee
    networks.set_trainer(conn, owner, north["id"], b["id"], svc.new_key())
    assert svc.get_user(conn, a["id"])["role"] == "employee" and svc.get_user(conn, b["id"])["role"] == "trainer"
    # the trainer role taken away frees the network
    svc.set_role(conn, owner, b["id"], "employee", svc.new_key())
    assert networks.get(conn, north["id"])["trainer_id"] is None

    with pytest.raises(svc.ServiceError) as e:
        networks.archive(conn, owner, south["id"], svc.new_key())
    assert e.value.code == "not_empty"
    networks.set_member(conn, owner, c["id"], None, svc.new_key())
    assert networks.archive(conn, owner, south["id"], svc.new_key())["status"] == "archived"


def test_shift_events_reach_the_network_trainer(conn, make_user):
    from app import networks

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer, u, other = make_user(first="Тренер"), make_user(first="Свой"), make_user(first="Чужой")
    net = networks.create(conn, owner, "Север", svc.new_key())
    networks.set_trainer(conn, owner, net["id"], trainer["id"], svc.new_key())
    networks.set_member(conn, owner, u["id"], net["id"], svc.new_key())
    s1 = shift(conn, svc.get_user(conn, u["id"]))
    s2 = shift(conn, other)
    n = FakeNotifier(conn)
    run(n.shift_event("SHIFT_CREATED", s1["id"]))
    run(n.shift_event("SHIFT_CREATED", s2["id"]))
    to_trainer = [t for chat, t in n.sent if chat == trainer["telegram_id"]]
    assert len(to_trainer) == 1 and "Свой" in to_trainer[0]


def test_migration_creates_networks_for_existing_ids(tmp_path):
    import sqlite3

    from app import db

    old = sqlite3.connect(tmp_path / "old.db", isolation_level=None)
    old.row_factory = sqlite3.Row
    old.executescript(db.SCHEMA)
    saved = db.MIGRATIONS[:]
    try:
        db.MIGRATIONS[:] = saved[:8]
        db.migrate(old)
        for tg, role in ((1, "trainer"), (2, "employee")):
            old.execute("INSERT INTO users (telegram_id, first_name, last_name, status, role, network_id, created_at) "
                        "VALUES (?, 'a', 'b', 'approved', ?, 3, 'x')", (tg, role))
        db.MIGRATIONS[:] = saved
        db.migrate(old)
    finally:
        db.MIGRATIONS[:] = saved
    row = old.execute("SELECT id, name, trainer_id FROM networks").fetchone()
    assert tuple(row) == (3, "Сеть 3", 1)


# ---------------------------------------------------------------- SBP profile and payouts (ТЗ модули 9–10)


def test_payment_profile_validation_and_audit(conn, make_user):
    from app import payouts

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    u, other = make_user(), make_user(first="Чужой")
    with pytest.raises(svc.ServiceError) as e:
        payouts.save_profile(conn, u, u["id"], "123", "Сбербанк", "Иван Петров", svc.new_key())
    assert e.value.code == "bad_phone"
    with pytest.raises(svc.ServiceError) as e:
        payouts.save_profile(conn, u, u["id"], "8 900 123 45 67", "Банк Луны", "Иван Петров", svc.new_key())
    assert e.value.code == "bad_bank"
    p = payouts.save_profile(conn, u, u["id"], "8 (900) 123-45-67", "Сбербанк", "Иван Петров", svc.new_key())
    assert p["phone"] == "+79001234567"
    with pytest.raises(svc.ServiceError) as e:  # nobody else but owners changes it
        payouts.save_profile(conn, other, u["id"], "+79000000000", "ВТБ", "Чужой Человек", svc.new_key())
    assert e.value.status == 403
    payouts.save_profile(conn, owner, u["id"], "+7 900 765 43 21", "Т-Банк", "Иван Петров", svc.new_key())
    log = payouts.profile_history(conn, u["id"])
    assert len(log) == 2 and '"Сбербанк"' in log[0]["old_value"]
    assert payouts.public_profile(payouts.get_profile(conn, u["id"]), False)["phone"] == "+7 *** ***-43-21"
    assert payouts.can_see(owner, u) and not payouts.can_see(other, u)


def test_payout_calculation_and_commit(conn, make_user):
    from app import finance, networks, payouts

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer, a, b = make_user(first="Тренер"), make_user(first="Аня"), make_user(first="Боря")
    net = networks.create(conn, owner, "Север", svc.new_key())
    networks.set_trainer(conn, owner, net["id"], trainer["id"], svc.new_key())
    networks.set_member(conn, owner, a["id"], net["id"], svc.new_key())
    a, b = svc.get_user(conn, a["id"]), svc.get_user(conn, b["id"])
    payouts.save_profile(conn, a, a["id"], "+79001112233", "Сбербанк", "Анна Аня", svc.new_key())

    s1 = shift(conn, a)                                         # 25 планом
    s2 = shift(conn, a, THU + timedelta(days=1), {**P, "volume": 10})
    approved(conn, a, s1, 30)                                   # approved, 30 × 120
    approved(conn, a, s2, 12, now=at(THU + timedelta(days=1), "19:30"))  # approved, lazy: 12 × 100
    sb = shift(conn, b)
    report(conn, b, sb, 10)                                     # below plan, not approved → not paid
    finance.create(conn, owner, a["id"], "fine", None, "Опоздание", svc.new_key(), event_date=THU, units=4)  # 4 × 120
    finance.create(conn, owner, a["id"], "bonus", 300, "Помощь", svc.new_key(), event_date=THU)
    finance.create(conn, owner, a["id"], "deduction", 200, "Форма", svc.new_key())
    finance.create(conn, owner, b["id"], "fine", None, "Поломка", svc.new_key(), event_date=THU, units=17)  # 17 × 120

    rates = payouts.Rates.from_dict(None)
    rows = {r["name"].split()[0]: r for r in payouts.calculate(conn, TODAY, THU + timedelta(days=3), rates)}
    ra = rows["Аня"]
    assert ra["personal_volume"] == 42 and ra["base_calc"] == 30 * 120 + 12 * 100
    assert (ra["bonuses"], ra["fines"], ra["deductions"]) == (300, 480, 200)
    assert ra["total"] == 4800 + 300 - 480 - 200 and ra["bank"] == "Сбербанк"
    assert rows["Тренер"]["team_volume"] == 42 and rows["Тренер"]["base_calc"] == 42 * 10
    rb = rows["Боря"]  # no approved work: the fine doesn't fit, the debt is kept, nothing negative
    assert rb["total"] == 0 and rb["fines"] == 0 and rb["carried_debt"] == 2040

    # a custom commission and a base override for a bonus
    custom = payouts.Rates.from_dict({"trainer_commission": 20})
    assert {r["name"].split()[0]: r for r in payouts.calculate(conn, TODAY, THU + timedelta(days=3), custom)}["Тренер"]["base_calc"] == 840

    with pytest.raises(svc.ServiceError) as e:
        payouts.commit(conn, a, "x", TODAY, THU, rates, None, [{"user_id": a["id"]}], [], [], svc.new_key())
    assert e.value.status == 403
    batch = payouts.commit(conn, owner, "Зарплата 12–18", TODAY, THU + timedelta(days=3), rates, None,
                           [{"user_id": a["id"]}, {"user_id": b["id"]}, {"user_id": trainer["id"], "base_override": 1000}],
                           ["proof-1"], [], svc.new_key())
    items = {i["name"].split()[0]: i for i in batch["items"]}
    assert items["Аня"]["total"] == 4420 and items["Тренер"]["base"] == 1000 and items["Боря"]["carried_debt"] == 2040
    assert batch["total"] == 4420 + 1000 and batch["proof_media"] == ["proof-1"]
    assert finance.balance(conn, a["id"])["balance"] == 0  # fine, bonus and deduction settled
    assert [t["status"] for t in finance.journal(conn, user_id=a["id"])] == ["offset", "offset", "offset"]
    assert finance.balance(conn, b["id"])["fines"] == 2040  # the debt stays for the next payout
    text = payouts.bank_text(batch)
    assert "Сбербанк:" in text and "Анна Аня — +7 900 111-22-33 — 4420 ₽" in text and "Всего к выплате: 5420 ₽" in text

    from app import export
    wb = load_workbook(io.BytesIO(export.payout_xlsx(batch)))
    assert wb.sheetnames == ["Реестр", "По банкам"]


def test_rates_come_from_settings(conn):
    from app import pricing

    app_settings.update(conn, {"rate_base": 150, "rate_lazy": 90, "lazy_threshold": 15})
    try:
        assert pricing.expected_pay(20) == 3000 and pricing.expected_pay(10) == 900
        assert pricing.pricing_info() == {"base_rate": 150, "lazy_rate": 90, "lazy_threshold": 15}
    finally:
        app_settings.update(conn, {"rate_base": 120, "rate_lazy": 100, "lazy_threshold": 20})


# ---------------------------------------------------------------- broadcasts with read receipts (ТЗ модуль 6)


def test_broadcast_audiences_by_network_and_people(conn, make_user):
    from app import networks

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer, a, b, c = make_user(first="Тренер"), make_user(first="Аня"), make_user(first="Боря", working=False), make_user(first="Вика")
    net = networks.create(conn, owner, "Север", svc.new_key())
    networks.set_trainer(conn, owner, net["id"], trainer["id"], svc.new_key())
    networks.set_member(conn, owner, a["id"], net["id"], svc.new_key())
    networks.set_member(conn, owner, b["id"], net["id"], svc.new_key())
    names = lambda rows: sorted(r["first_name"] for r in rows)  # noqa: E731
    assert names(svc.broadcast_recipients(conn, [], network_id=net["id"])) == ["Аня", "Боря", "Тренер"]
    assert names(svc.broadcast_recipients(conn, [], user_ids=[c["id"]])) == ["Вика"]
    assert names(svc.broadcast_recipients(conn, ["not_working"], user_ids=[c["id"]])) == ["Боря", "Вика", "Саня"]
    with pytest.raises(svc.ServiceError):
        svc.broadcast_recipients(conn, [])
    # a trainer reaches only their network whatever they choose
    t = svc.get_user(conn, trainer["id"])
    assert names(svc.broadcast_recipients(conn, ["working", "not_working"], actor=t)) == ["Аня", "Боря"]

    # the bug «6 вместо 1»: a trainer ticks one person — only that person gets it, not the whole network
    row, rec = svc.create_broadcast(conn, t, [], "Аня, зайди", user_ids=[a["id"]])
    assert names(rec) == ["Аня"] and row["recipients"] == 1
    assert names(svc.broadcast_recipients(conn, [], user_ids=[a["id"], c["id"]], actor=t)) == ["Аня"]  # not their network


def test_broadcast_only_to_those_who_have_not_marked_tomorrow(conn, make_user):
    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b, c, d = (make_user(first=n) for n in ("Аня", "Боря", "Вика", "Гоша"))
    tomorrow = TODAY + timedelta(days=1)
    shift(conn, a, tomorrow)  # going
    svc.pre_declare_off(conn, b, tomorrow, svc.new_key(), today=TODAY)  # a day off is marked too
    gone = shift(conn, c, tomorrow)
    svc.cancel_shift(conn, c, gone["id"], svc.new_key(), today=TODAY)  # «Беру выходной»: the day is marked as off
    assert svc.marked_on(conn, tomorrow) == {a["id"], b["id"], c["id"]}
    offset = 0 if tomorrow.isocalendar()[1] == TODAY.isocalendar()[1] else 1
    day = next(d for d in svc.week_view(conn, a, offset, today=TODAY)["days"] if d["date"] == tomorrow.isoformat())
    assert sorted(o["first_name"] for o in day["colleagues_off"]) == ["Боря", "Вика"]  # colleagues see it as off
    everyone = [a["id"], b["id"], c["id"], d["id"]]
    row, rec = svc.create_broadcast(conn, owner, [], "Отметьте график на завтра", user_ids=everyone, only_unmarked=tomorrow)
    assert sorted(u["first_name"] for u in rec) == ["Гоша"] and "unmarked_tomorrow" in row["audience"]
    with pytest.raises(svc.ServiceError) as e:
        svc.create_broadcast(conn, owner, [], "Привет", user_ids=[a["id"]], only_unmarked=tomorrow)
    assert e.value.code == "no_recipients"


def test_read_receipts_and_reminders(conn, make_user):
    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b = make_user(first="Аня"), make_user(first="Боря")
    row, recipients = svc.create_broadcast(conn, owner, [], "<b>Важно</b>", True, user_ids=[a["id"], b["id"]],
                                           require_ack=True)
    n = FakeNotifier(conn)
    assert run(n.broadcast(row["id"], recipients, row["html"], True, None, True)) == 2
    assert all(m is not None and m.inline_keyboard[0][0].callback_data == f"ack:{row['id']}" for _, m in n.markups)
    assert svc.acknowledge(conn, row["id"], a["telegram_id"]) is True
    assert svc.acknowledge(conn, row["id"], a["telegram_id"]) is False  # once
    receipts = {d["first_name"]: d for d in svc.broadcast_deliveries(conn, row["id"])}
    assert receipts["Аня"]["read_at"] and receipts["Боря"]["read_at"] is None
    assert svc.list_broadcasts(conn)[0]["read"] == 1
    n.sent.clear()
    assert run(n.broadcast_remind(row["id"])) == 1  # only the one who hasn't confirmed
    assert n.sent[0][0] == b["telegram_id"] and "Напоминание" in n.sent[0][1]
    assert svc.broadcast_deliveries(conn, row["id"])[0]["reminders"] == 1


# ---------------------------------------------------------------- goals (ТЗ модуль 5)


def test_goals_progress_and_completion(conn, make_user):
    from app import goals, networks

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer, a, b, c = make_user(first="Тренер"), make_user(first="Аня"), make_user(first="Боря"), make_user(first="Вика")
    net = networks.create(conn, owner, "Север", svc.new_key())
    networks.set_trainer(conn, owner, net["id"], trainer["id"], svc.new_key())
    for u in (a, b):
        networks.set_member(conn, owner, u["id"], net["id"], svc.new_key())
    trainer, a, b = (svc.get_user(conn, u["id"]) for u in (trainer, a, b))

    with pytest.raises(svc.ServiceError):  # company-wide goals are for owners and admins
        goals.create(conn, trainer, title="Все", kind="team", target=10, date_from=TODAY, date_to=THU, key=svc.new_key())
    with pytest.raises(svc.ServiceError):  # not their person
        goals.create(conn, trainer, title="Лично", kind="personal", target=10, date_from=TODAY, date_to=THU,
                     user_id=c["id"], key=svc.new_key())
    personal = goals.create(conn, trainer, title="Лично Ане", kind="personal", target=25, date_from=TODAY, date_to=THU,
                            user_id=a["id"], key=svc.new_key())
    shared = goals.create(conn, trainer, title="Северу 50", kind="network", target=50, date_from=TODAY, date_to=THU,
                          network_id=net["id"], key=svc.new_key())
    team = goals.create(conn, owner, title="Каждый по 20", kind="team", mode="each", target=20, date_from=TODAY,
                        date_to=THU, key=svc.new_key())

    assert [g["title"] for g in goals.for_user(conn, a, TODAY)] == ["Лично Ане", "Северу 50", "Каждый по 20"]
    assert [g["title"] for g in goals.for_user(conn, c, TODAY)] == ["Каждый по 20"]
    assert goals.view(conn, personal, a)["joined"]  # personal goals join automatically
    with pytest.raises(svc.ServiceError):
        goals.join(conn, c, shared["id"], TODAY)  # not in the network
    for u in (a, b):
        goals.join(conn, u, shared["id"], TODAY)
    goals.join(conn, a, team["id"], TODAY)
    assert goals.join(conn, a, team["id"], TODAY)[1] is False

    sa, sb = shift(conn, a), shift(conn, b)
    approved(conn, a, sa, 30)  # approved: counts in every goal at once
    report(conn, b, sb, 10)  # below plan, waits for review → not counted yet
    v = goals.view(conn, goals.get(conn, shared["id"]), a)
    assert v["progress"] == 30 and v["my_volume"] == 30

    n = FakeNotifier(conn)
    run(jobs.tick(conn, n, at(THU, "20:01")))
    assert goals.get(conn, personal["id"])["status"] == "completed"
    assert goals.get(conn, shared["id"])["status"] == "active"
    assert any("Цель выполнена" in t and "Лично Ане" in t for chat, t in n.sent if chat == a["telegram_id"])
    team_view = goals.view(conn, goals.get(conn, team["id"]), a)
    assert team_view["done"] and team_view["participants"][0]["done"]

    reports.review(conn, owner, sb["id"], "approve", svc.new_key())  # 30 + 10 = 40 < 50
    assert goals.check_completions(conn, THU) == []
    sb2 = shift(conn, b, THU + timedelta(days=1))
    report(conn, b, sb2, 25, now=at(THU + timedelta(days=1), "19:30"))  # outside the period → doesn't count
    assert goals.check_completions(conn, THU) == []
    lifecycle.set_visibility(conn, trainer, "goal", shared["id"], True, svc.new_key())
    assert all(g["title"] != "Северу 50" for g in goals.for_user(conn, a, TODAY))


# ---------------------------------------------------------------- live events (ТЗ модуль 7)


def _event(conn, actor, **kw):
    from app import events

    base = dict(title="Гонка", mode="first_to_reach", target=30, max_winners=1, max_attempts=2,
                starts_at=at(THU, "10:00"), ends_at=at(THU, "18:00"), key=svc.new_key())
    return events.create(conn, actor, **{**base, **kw})


def test_first_to_reach_race(conn, make_user):
    from app import events

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b, c = make_user(first="Аня"), make_user(first="Боря"), make_user(first="Вика")
    ev = _event(conn, owner, max_winners=2)
    noon = at(THU, "12:00")
    with pytest.raises(svc.ServiceError) as e:
        events.submit(conn, a, ev["id"], 31, "p", svc.new_key(), now=at(THU, "09:00"))
    assert e.value.code == "closed"
    with pytest.raises(svc.ServiceError) as e:
        events.submit(conn, a, ev["id"], 20, "p", svc.new_key(), now=noon)
    assert e.value.code == "below_target"
    ea, _ = events.submit(conn, a, ev["id"], 31, "p", svc.new_key(), now=noon)
    with pytest.raises(svc.ServiceError) as e:  # one entry under review at a time
        events.submit(conn, a, ev["id"], 32, "p", svc.new_key(), now=noon)
    assert e.value.code == "pending"
    board = events.view(conn, ev, a, noon)["leaderboard"]
    assert board[0]["name"].startswith("Аня") and board[0]["score"] == 31 and not board[0]["winner"]  # soft score

    eb, _ = events.submit(conn, b, ev["id"], 30, "p", svc.new_key(), now=noon)
    with pytest.raises(svc.ServiceError) as e:
        events.review(conn, owner, eb["id"], False, "", svc.new_key())
    assert e.value.code == "reason_required"
    events.review(conn, owner, eb["id"], False, "Скрин не тот", svc.new_key())
    eb2, _ = events.submit(conn, b, ev["id"], 30, "p", svc.new_key(), now=noon)
    events.review(conn, owner, eb2["id"], False, "Опять не тот", svc.new_key())
    vb = events.view(conn, events.get(conn, ev["id"]), b, noon)
    assert vb["my"]["attempts_left"] == 0 and not vb["my"]["can_submit"]
    with pytest.raises(svc.ServiceError) as e:
        events.submit(conn, b, ev["id"], 40, "p", svc.new_key(), now=noon)
    assert e.value.code == "no_attempts"

    first = events.review(conn, owner, ea["id"], True, None, svc.new_key())
    assert first.changed and not first.finished
    ec, _ = events.submit(conn, c, ev["id"], 35, "p", svc.new_key(), now=at(THU, "13:00"))
    second = events.review(conn, owner, ec["id"], True, None, svc.new_key())
    assert second.finished and events.get(conn, ev["id"])["status"] == "finished"
    board = events.leaderboard(conn, events.get(conn, ev["id"]))
    assert [r["name"].split()[0] for r in board if r["winner"]] == ["Аня", "Вика"]  # 👑 in confirmation order
    n = FakeNotifier(conn)
    run(n.event_reviewed(second, owner))
    assert any("завершено" in t for chat, t in n.sent if chat == b["telegram_id"])


def test_accumulative_scoreboard_and_finish(conn, make_user):
    from app import events

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b = make_user(first="Аня"), make_user(first="Боря")
    ev = _event(conn, owner, title="Кто больше", mode="accumulative", target=None, max_winners=1)
    e1, _ = events.submit(conn, a, ev["id"], 10, "p", svc.new_key(), now=at(THU, "11:00"))
    events.review(conn, owner, e1["id"], True, None, svc.new_key())
    e2, _ = events.submit(conn, b, ev["id"], 12, "p", svc.new_key(), now=at(THU, "12:00"))
    e3, _ = events.submit(conn, a, ev["id"], 5, "p", svc.new_key(), now=at(THU, "12:30"))
    board = events.view(conn, ev, None, at(THU, "13:00"))["leaderboard"]
    assert [(r["name"].split()[0], r["score"], r["confirmed"]) for r in board] == [("Аня", 15, 10), ("Боря", 12, 0)]
    events.review(conn, owner, e2["id"], True, None, svc.new_key())
    events.review(conn, owner, e3["id"], False, "Не видно итога", svc.new_key())
    with pytest.raises(svc.ServiceError) as e:  # after the timer: closed for new entries
        events.submit(conn, a, ev["id"], 50, "p", svc.new_key(), now=at(THU, "18:05"))
    assert e.value.code == "closed"
    n = FakeNotifier(conn)
    run(jobs.tick(conn, n, at(THU, "18:10")))  # still in the 15-minute buffer
    assert events.get(conn, ev["id"])["status"] == "active"
    run(jobs.tick(conn, n, at(THU, "18:16")))
    assert events.get(conn, ev["id"])["status"] == "finished"
    board = events.leaderboard(conn, events.get(conn, ev["id"]))
    assert [r["name"].split()[0] for r in board if r["winner"]] == ["Боря"]  # prizes by confirmed score: 12 > 10
    assert any("Победители" in t and "Боря" in t for chat, t in n.sent if chat == a["telegram_id"])
    # the creator is also an eligible participant: one message, not two; a later tick doesn't repeat it
    run(jobs.tick(conn, n, at(THU, "18:17")))
    assert [chat for chat, t in n.sent if "завершено" in t].count(999) == 1
    assert [chat for chat, t in n.sent if "завершено" in t].count(a["telegram_id"]) == 1


def test_event_network_scope(conn, make_user):
    from app import events, networks

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer, a, other = make_user(first="Тренер"), make_user(first="Аня"), make_user(first="Чужой")
    net = networks.create(conn, owner, "Север", svc.new_key())
    networks.set_trainer(conn, owner, net["id"], trainer["id"], svc.new_key())
    networks.set_member(conn, owner, a["id"], net["id"], svc.new_key())
    trainer, a = svc.get_user(conn, trainer["id"]), svc.get_user(conn, a["id"])
    with pytest.raises(svc.ServiceError):
        _event(conn, trainer)  # a trainer runs events only for their network
    ev = _event(conn, trainer, network_id=net["id"])
    assert events.eligible(ev, a) and not events.eligible(ev, other)
    with pytest.raises(svc.ServiceError) as e:
        events.submit(conn, other, ev["id"], 40, "p", svc.new_key(), now=at(THU, "12:00"))
    assert e.value.status == 404


# ---------------------------------------------------------------- broadcasts: video, templates, «Ответить»


def test_broadcast_video_reply_button_and_replies(conn, make_user):
    from app import broadcasts, media
    from app.notify import reply_text

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a, b = make_user(first="Аня"), make_user(first="Боря")
    row, rec = svc.create_broadcast(conn, owner, [], "Смотрите видео", user_ids=[a["id"], b["id"]], require_ack=True,
                                    media="video-file", media_kind="video", allow_reply=True)
    n = FakeNotifier(conn)
    assert run(n.broadcast(row["id"], rec, row["html"], True, "video-file", True, "video", True)) == 2
    buttons = [btn.callback_data for btn in n.markups[0][1].inline_keyboard[0]]
    assert buttons == [f"ack:{row['id']}", f"brr:{row['id']}"] and n.kinds == ["video", "video"]

    with pytest.raises(svc.ServiceError) as e:  # nothing to send
        broadcasts.add_reply(conn, row["id"], a, "  ", [])
    assert e.value.code == "empty"
    outsider = make_user(first="Чужой")
    with pytest.raises(svc.ServiceError):  # it wasn't sent to them
        broadcasts.add_reply(conn, row["id"], outsider, "Привет", [])
    b_row, reply = broadcasts.add_reply(conn, row["id"], a, "Посмотрела, всё понятно",
                                        [{"kind": "photo", "file_id": "p1"}, {"kind": "video", "file_id": "v1"}])
    assert reply["media"][1]["kind"] == "video"
    text = reply_text(b_row, svc.get_user(conn, a["id"]), reply["text"], "Север")
    assert f'tg://user?id={a["telegram_id"]}' in text and "Аня" in text and "сеть «Север»" in text and "всё понятно" in text
    # a reply counts as read; the history shows who replied
    got = {r["user_id"]: r for r in svc.broadcast_deliveries(conn, row["id"])}
    assert got[a["id"]]["read_at"] and not got[b["id"]]["read_at"]
    assert svc.list_broadcasts(conn)[0]["replied"] == 1
    assert broadcasts.replies_by_user(conn, row["id"])[a["id"]][0] | {"created_at": None} == {
        "text": "Посмотрела, всё понятно", "photos": 1, "videos": 1, "created_at": None}
    # replies off: no button, no reply
    quiet, rec2 = svc.create_broadcast(conn, owner, [], "Без ответа", user_ids=[a["id"]])
    with pytest.raises(svc.ServiceError) as e:
        broadcasts.add_reply(conn, quiet["id"], a, "А можно?", [])
    assert e.value.code == "no_reply"
    assert media.sniff_video(b"\x00\x00\x00\x20ftypisom") == "mp4" and media.sniff_video(b"\x00\x00\x00\x14ftypqt  ") == "mov"
    assert media.sniff_video(b"GIF89a....") is None


def test_broadcast_templates(conn, make_user):
    from app import broadcasts, networks

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer = make_user(first="Тренер")
    net = networks.create(conn, owner, "Север", svc.new_key())
    networks.set_trainer(conn, owner, net["id"], trainer["id"], svc.new_key())
    trainer = svc.get_user(conn, trainer["id"])
    t = broadcasts.save_template(conn, owner, "  Сбор  ", "<b>Завтра</b> сбор в 9:00<script>x</script>")
    assert t["title"] == "Сбор" and "<b>Завтра</b>" in t["html"] and "script" not in t["html"]
    mine = broadcasts.save_template(conn, trainer, "Мой", "Привет, Север")
    assert [x["title"] for x in broadcasts.templates(conn)] == ["Мой", "Сбор"]  # everyone may use every template
    with pytest.raises(svc.ServiceError) as e:  # but changes only their own (owners and admins — any)
        broadcasts.save_template(conn, trainer, "Сбор", "чужой", t["id"])
    assert e.value.status == 403
    assert broadcasts.save_template(conn, owner, "Мой", "Привет, Север!", mine["id"])["html"] == "Привет, Север!"
    with pytest.raises(svc.ServiceError):
        broadcasts.save_template(conn, owner, "", "текст")
    broadcasts.delete_template(conn, trainer, mine["id"])
    assert [x["title"] for x in broadcasts.templates(conn)] == ["Сбор"]


def test_sent_broadcast_media_opens_in_full(conn, make_user):
    """«Отправленные»: the photo or video of a sent broadcast opens in the app (videos — from the kept copy)."""
    from fastapi.testclient import TestClient

    from app import media
    from app.api import create_app
    from app.auth import sign_init_data

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    a = make_user(first="Аня")
    clip = b"\x00\x00\x00\x20ftypisom" + b"\x00" * 64
    media.keep_video("tg-video-1", clip, "mp4")  # what store_video keeps after the upload to Telegram
    row, _ = svc.create_broadcast(conn, owner, [], "<b>Длинный</b> текст", user_ids=[a["id"]], media="tg-video-1",
                                  media_kind="video")
    client = TestClient(create_app(conn, FakeNotifier(conn)))
    headers = {"X-Telegram-Init-Data": sign_init_data({"id": 999, "first_name": "Саня"}, settings.bot_token)}
    item = client.get("/api/admin/broadcasts", headers=headers).json()["items"][0]
    assert item["id"] == row["id"] and item["media_url"].startswith(f"/api/media/broadcast/{row['id']}?")
    got = client.get(item["media_url"])
    assert got.status_code == 200 and got.headers["content-type"] == "video/mp4" and got.content == clip
    part = client.get(item["media_url"], headers={"Range": "bytes=0-7"})  # the player asks for pieces
    assert part.status_code == 206 and part.content == clip[:8]
    assert client.get(item["media_url"].replace("sig=", "sig=0")).status_code == 403
    plain, _ = svc.create_broadcast(conn, owner, [], "Без медиа", user_ids=[a["id"]])
    assert client.get("/api/admin/broadcasts", headers=headers).json()["items"][0]["media_url"] is None


def test_forgive_fines_trainer_and_admin_reason_optional(conn, make_user):
    """«Простить»: a trainer — their network's fines, an admin — anyone's; the reason is optional."""
    from app import finance

    owner = svc.ensure_admin(conn, 999, "Саня", "Владелец", "boss")
    trainer = _staff(conn, make_user, "trainer", "Тренер", network=1)
    admin = _staff(conn, make_user, "admin", "Админ")
    u = make_user(first="Иван")
    conn.execute("UPDATE users SET network_id = 1 WHERE id = ?", (u["id"],))
    stranger = make_user(first="Чужой")
    mine = finance.create(conn, admin, u["id"], "fine", None, "Опоздал", svc.new_key(), units=5).tx
    theirs = finance.create(conn, admin, stranger["id"], "fine", None, "Опоздал", svc.new_key(), units=3).tx
    bonus = finance.create(conn, admin, u["id"], "bonus", 200, "Молодец", svc.new_key()).tx
    with pytest.raises(svc.ServiceError) as e:  # not their network
        finance.forgive(conn, trainer, theirs["id"], None, svc.new_key())
    assert e.value.status == 403
    with pytest.raises(svc.ServiceError) as e:
        finance.forgive(conn, admin, bonus["id"], None, svc.new_key())
    assert e.value.code == "not_a_fine"
    done = finance.forgive(conn, trainer, mine["id"], None, svc.new_key())  # no reason
    assert done.tx["status"] == "cancelled" and done.tx["decision_comment"] == "Прощён"
    assert finance.balance(conn, u["id"])["balance"] == 200
    with pytest.raises(svc.ServiceError):  # once
        finance.forgive(conn, admin, mine["id"], None, svc.new_key())
    with_reason = finance.forgive(conn, admin, theirs["id"], "Больничный", svc.new_key())
    assert with_reason.tx["decision_comment"] == "Прощён: Больничный"
    # a trainer's request waiting for approval can be forgiven too; the employee isn't told about it
    req = finance.create(conn, trainer, u["id"], "fine", None, "Спорно", svc.new_key(), units=4).tx
    n = FakeNotifier(conn)
    run(n.tx_forgiven(finance.forgive(conn, trainer, req["id"], None, svc.new_key()).tx, trainer,
                      "pending_admin_approval", None))
    assert not any(chat == u["telegram_id"] for chat, _ in n.sent)
    n.sent.clear()
    fresh = finance.create(conn, owner, u["id"], "fine", None, "Мусор", svc.new_key(), units=1).tx
    run(n.tx_forgiven(finance.forgive(conn, admin, fresh["id"], "Убрал сам", svc.new_key()).tx, admin, "active", "Убрал сам"))
    assert any(chat == u["telegram_id"] and "прощён" in t and "Убрал сам" in t for chat, t in n.sent)

