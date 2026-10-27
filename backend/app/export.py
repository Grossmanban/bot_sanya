"""Excel export of shifts, per-person totals and penalties for a period."""

import io
import sqlite3
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import finance, pricing
from .services import ACTIVE, full_name, shortfall

STATUS_RU = {
    "active": "Выход",
    "restored": "Выход (восстановлен)",
    "cancelled": "Отменён",
    "finally_cancelled": "Отменён со штрафом",
    "awaiting_approval": "Ждёт решения",
    "rejected_by_admin": "Отклонён",
    "pre_declared_off": "Выходной",
}
HEADER_FILL = PatternFill("solid", fgColor="161A33")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    return ws


def build_xlsx(conn: sqlite3.Connection, date_from: date, date_to: date, user_id: int | None = None,
               only: set[int] | None = None) -> bytes:
    """only: restrict to these user ids (a trainer's network)."""
    params: list = [date_from.isoformat(), date_to.isoformat()]
    where = "s.shift_date BETWEEN ? AND ? AND s.status NOT IN ('withdrawn','removed')"
    if user_id:
        where += " AND s.user_id = ?"
        params.append(user_id)
    if only is not None:
        where += f" AND s.user_id IN ({','.join('?' * len(only)) or 'NULL'})"
        params.extend(sorted(only))
    shifts = [
        dict(r)
        for r in conn.execute(
            f"""SELECT s.*, u.first_name, u.last_name, u.username FROM shifts s JOIN users u ON u.id = s.user_id
                WHERE {where} ORDER BY s.shift_date, s.start_time IS NULL, s.start_time, u.first_name""",
            params,
        )
    ]

    wb = Workbook()
    from .rates import Book

    book = Book.load(conn)

    def own(s: dict) -> dict:  # the person's rates: personal / network / common
        return book.user_rates(s["user_id"])

    totals: dict[int, dict] = {}
    for s in shifts:
        t = totals.setdefault(
            s["user_id"],
            {"name": full_name(s), "username": s["username"] or "", "shifts": 0, "plan": 0, "fact": 0, "no_fact": 0,
             "plan_pay": 0, "fact_pay": 0, "cancelled": 0, "penalties": 0, "off": 0, "short": 0, "short_shifts": 0},
        )
        if s["status"] in ACTIVE:
            t["shifts"] += 1
            t["plan"] += s["volume"]
            t["plan_pay"] += pricing.expected_pay(s["volume"], own(s))
            if s["fact_volume"] is not None:
                t["fact"] += s["fact_volume"]
                t["fact_pay"] += pricing.expected_pay(s["fact_volume"], own(s))
            else:
                t["no_fact"] += 1
        elif s["status"] in ("cancelled", "finally_cancelled"):
            t["cancelled"] += 1
        elif s["status"] == "pre_declared_off":
            t["off"] += 1
        if shortfall(s):
            t["short"] += shortfall(s)
            t["short_shifts"] += 1
    # money comes from the ledger (ТЗ модуль 4): fines and bonuses that took effect in the period
    money = [t for t in finance.journal(conn, user_id=user_id, only=only, date_from=date_from, date_to=date_to, limit=100_000)]
    for tx in money:
        if tx["status"] not in finance.HAPPENED:
            continue
        user = conn.execute("SELECT * FROM users WHERE id = ?", (tx["user_id"],)).fetchone()
        t = totals.setdefault(
            tx["user_id"],
            {"name": full_name(dict(user)), "username": user["username"] or "", "shifts": 0, "plan": 0, "fact": 0,
             "no_fact": 0, "plan_pay": 0, "fact_pay": 0, "cancelled": 0, "penalties": 0, "off": 0, "short": 0,
             "short_shifts": 0},
        )
        key = {"fine": "penalties", "bonus": "bonuses", "deduction": "deductions"}.get(tx["type"])
        if key:
            t[key] = t.get(key, 0) + tx["amount"]
    total_rows = [
        [t["name"], t["username"], t["shifts"], t["plan"], t["fact"], t["no_fact"], t["short_shifts"], t["short"],
         t["plan_pay"], t["fact_pay"], t["cancelled"], t["off"], t["penalties"], t.get("bonuses", 0),
         t.get("deductions", 0), t["fact_pay"] - t["penalties"] + t.get("bonuses", 0) - t.get("deductions", 0)]
        for t in sorted(totals.values(), key=lambda x: x["name"])
    ]
    if total_rows:
        total_rows.append(["Итого", "", *[sum(r[i] for r in total_rows) for i in range(2, 16)]])
    ws = _sheet(
        wb, "Итоги",
        ["Сотрудник", "Username", "Выходов", "План, шт.", "Факт, шт.", "Без факта", "Не по норме, смен",
         "Недобор, шт.", "План, ₽", "Факт, ₽", "Отмен", "Выходных", "Штрафы, ₽", "Бонусы, ₽", "Вычеты, ₽",
         "Итог (факт − штрафы + бонусы − вычеты), ₽"],
        total_rows, [24, 16, 10, 11, 11, 11, 18, 13, 12, 12, 9, 11, 12, 12, 12, 30], first=True,
    )
    if total_rows:
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    _sheet(
        wb, "Смены",
        ["Дата", "Сотрудник", "Статус", "Начало", "Конец", "Точка", "План, шт.", "Ставка, ₽", "План, ₽",
         "Факт, шт.", "Факт, ₽", "Недобор, шт.", "Штраф за норму, ₽", "Штраф за отмену, ₽", "Срочная"],
        [
            [
                date.fromisoformat(s["shift_date"]).strftime("%d.%m.%Y"), full_name(s), STATUS_RU.get(s["status"], s["status"]),
                s["start_time"] or "", s["end_time"] or "", s["location"] or "", s["volume"] or "",
                pricing.calculate_rate(s["volume"], own(s)) if s["volume"] else "",
                pricing.expected_pay(s["volume"], own(s)) if s["volume"] and s["status"] in ACTIVE else "",
                s["fact_volume"] if s["fact_volume"] is not None else "",
                pricing.expected_pay(s["fact_volume"], own(s)) if s["fact_volume"] is not None else "",
                shortfall(s) or "",
                s["shortfall_penalty"] if s["shortfall_penalty"] is not None else ("без штрафа" if s["shortfall_status"] == "waived" else ""),
                s["penalty_amount"] if s["penalty_amount"] is not None else "",
                "да" if s["is_urgent"] else "",
            ]
            for s in shifts
        ],
        [12, 24, 20, 9, 9, 16, 10, 10, 10, 10, 10, 12, 18, 18, 9],
    )

    _sheet(
        wb, "Штрафы и бонусы",
        ["Дата", "Сотрудник", "Тип", "Причина", "Сумма, ₽", "Статус", "Назначил", "Решение"],
        [
            [date.fromisoformat(tx["event_date"]).strftime("%d.%m.%Y"), tx["name"], finance.TYPE_TITLES[tx["type"]],
             tx["reason"], tx["amount"] if tx["amount"] is not None else "не назначен",
             finance.STATUS_TITLES[tx["status"]], tx["created_by"] or "", tx["decision_comment"] or ""]
            for tx in money
        ],
        [12, 24, 12, 40, 12, 18, 22, 28],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename(date_from: date, date_to: date) -> str:
    return f"grafik_{date_from:%d.%m}-{date_to:%d.%m.%Y}.xlsx"


ROLE_RU = {"owner": "Владелец", "admin": "Администратор", "trainer": "Тренер", "instructor": "Инструктор",
           "employee": "Сотрудник"}


def payout_xlsx(batch: dict) -> bytes:
    """Payout registry (ТЗ 11.3): one row per person, then a sheet grouped by bank."""
    wb = Workbook()
    rows = [
        [it["name"], ROLE_RU.get(it["role"], it["role"]), it["network"] or "", it["phone"] or "", it["bank"] or "",
         it["recipient"] or "", it["personal_volume"], it["team_volume"], it["base"], it["bonuses"], it["fines"],
         it["deductions"], it["total"], it["carried_debt"] or ""]
        for it in batch["items"]
    ]
    rows.append(["Итого", "", "", "", "", "", sum(r[6] for r in rows), sum(r[7] for r in rows), sum(r[8] for r in rows),
                 sum(r[9] for r in rows), sum(r[10] for r in rows), sum(r[11] for r in rows), batch["total"],
                 sum(r[13] or 0 for r in rows)])
    ws = _sheet(
        wb, "Реестр",
        ["ФИО", "Роль", "Сеть", "Телефон (СБП)", "Банк", "Получатель", "Объём, шт.", "Командный объём, шт.", "База, ₽",
         "Бонусы, ₽", "Штрафы, ₽", "Вычеты, ₽", "Итого к выплате, ₽", "Остаток долга, ₽"],
        rows, [24, 14, 12, 18, 16, 26, 11, 18, 11, 11, 11, 11, 18, 16], first=True,
    )
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    by_bank = sorted((it for it in batch["items"] if it["total"] > 0), key=lambda it: (it["bank"] or "яя", it["name"]))
    _sheet(wb, "По банкам", ["Банк", "Получатель", "Телефон (СБП)", "Сумма, ₽"],
           [[it["bank"] or "Реквизиты не указаны", it["recipient"] or it["name"], it["phone"] or "", it["total"]] for it in by_bank],
           [20, 28, 18, 12])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
